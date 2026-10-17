"""Turns source documents into text chunks that keep their provenance."""
from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md", ".markdown"}
MAX_FILE_MB = 50
CHUNK_WORDS = 120
CHUNK_OVERLAP = 30


class IngestionError(RuntimeError):
    pass


@dataclass
class Chunk:
    chunk_id: str
    source_id: str
    source: str
    text: str
    page: int | None = None
    heading: str | None = None

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class SourceDoc:
    source_id: str
    name: str
    kind: str
    pages: int | None
    blocks: list[tuple[str, int | None, str | None]]  # (text, page, heading)

    @property
    def chars(self) -> int:
        return sum(len(b[0]) for b in self.blocks)


def _clean(text: str) -> str:
    text = text.replace("Â­", "")
    text = re.sub(r"-\n(\w)", r"\1", text)  # re-join hyphenated line breaks
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def read_pdf(path: Path) -> tuple[list[tuple[str, int | None, str | None]], int]:
    import pymupdf

    blocks = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, 1):
            t = _clean(page.get_text("text"))
            if t:
                blocks.append((t, i, None))
        return blocks, doc.page_count


def read_docx(path: Path) -> list[tuple[str, int | None, str | None]]:
    import docx

    d = docx.Document(str(path))
    blocks, heading, buf = [], None, []
    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if (p.style is not None and p.style.name.lower().startswith("heading")) or p.style.name == "Title":
            if buf:
                blocks.append(("\n".join(buf), None, heading))
                buf = []
            heading = txt
        else:
            buf.append(txt)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                buf.append(" | ".join(cells))
    if buf:
        blocks.append(("\n".join(buf), None, heading))
    return blocks


def read_markdown(text: str) -> list[tuple[str, int | None, str | None]]:
    blocks, heading, buf = [], None, []
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}#{1,6}\s+(.*)", line)
        if m:
            if "".join(buf).strip():
                blocks.append((_clean("\n".join(buf)), None, heading))
            buf, heading = [], m.group(1).strip()
        else:
            buf.append(line)
    if "".join(buf).strip():
        blocks.append((_clean("\n".join(buf)), None, heading))
    return blocks


def load_source(path: str | Path, source_id: str) -> SourceDoc:
    path = Path(path)
    if not path.is_file():
        raise IngestionError(f"source file not found: {path.name}")
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise IngestionError(f"unsupported source type {ext!r} for {path.name}")
    if path.stat().st_size > MAX_FILE_MB * 1_000_000:
        raise IngestionError(f"{path.name} is larger than {MAX_FILE_MB} MB")
    pages = None
    try:
        if ext == ".pdf":
            blocks, pages = read_pdf(path)
        elif ext == ".docx":
            blocks = read_docx(path)
        elif ext in (".md", ".markdown"):
            blocks = read_markdown(path.read_text(encoding="utf-8-sig", errors="replace"))
        else:
            blocks = [(_clean(path.read_text(encoding="utf-8-sig", errors="replace")), None, None)]
    except IngestionError:
        raise
    except Exception as e:  # corrupt files
        raise IngestionError(f"could not read {path.name}: {e}") from e
    blocks = [b for b in blocks if b[0].strip()]
    if not blocks:
        raise IngestionError(f"no extractable text in {path.name} (scanned PDFs need OCR first)")
    return SourceDoc(source_id=source_id, name=path.name, kind=ext.lstrip("."), pages=pages, blocks=blocks)


def text_source(text: str, source_id: str, name: str = "brief_source_text") -> SourceDoc:
    blocks = read_markdown(text) or [(_clean(text), None, None)]
    return SourceDoc(source_id=source_id, name=name, kind="text", pages=None, blocks=blocks)


def chunk(doc: SourceDoc, words: int = CHUNK_WORDS, overlap: int = CHUNK_OVERLAP) -> list[Chunk]:
    out: list[Chunk] = []
    for text, page, heading in doc.blocks:
        # split on sentences, then pack into windows of ~`words` words
        sents = re.split(r"(?<=[.!?])\s+", text.replace("\n", " "))
        cur: list[str] = []
        # very long "sentences" (lists, tables, missing punctuation) are cut into word windows
        pieces = []
        for s in sents:
            w = s.split()
            pieces += [" ".join(w[i:i + words]) for i in range(0, len(w), words)] or [""]
        for s in pieces:
            sw = s.split()
            if cur and len(cur) + len(sw) > words:
                out.append(Chunk(f"{doc.source_id}-C{len(out) + 1:03d}", doc.source_id, doc.name, " ".join(cur),
                                 page, heading))
                cur = cur[-overlap:] if overlap else []
            cur.extend(sw)
        if cur:
            out.append(Chunk(f"{doc.source_id}-C{len(out) + 1:03d}", doc.source_id, doc.name, " ".join(cur), page,
                             heading))
    return out


