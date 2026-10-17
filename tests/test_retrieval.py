import docx
import pymupdf
import pytest

from retrieval.index import ChunkIndex
from retrieval.ingestion import IngestionError, chunk, load_source, text_source
from retrieval.query import numbers_supported, verify_claim


@pytest.fixture
def docs(tmp_path):
    md = tmp_path / "notes.md"
    md.write_text("# Bees\n\n## Pollination\nBees pollinate many flowering plants while collecting nectar.\n",
                  encoding="utf-8")
    txt = tmp_path / "plain.txt"
    txt.write_text("Solar panels convert sunlight into electricity using photovoltaic cells.", encoding="utf-8")
    d = docx.Document()
    d.add_heading("Volcanoes", level=1)
    d.add_paragraph("Magma that reaches the surface is called lava.")
    dx = tmp_path / "geo.docx"
    d.save(dx)
    pdf = tmp_path / "space.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Page one is an introduction.")
    doc.new_page().insert_text((72, 72), "The Moon orbits the Earth about once every 27 days.")
    doc.save(pdf)
    return md, txt, dx, pdf


def test_ingest_all_formats_with_metadata(docs):
    md, txt, dx, pdf = docs
    s_md = load_source(md, "S001")
    assert s_md.blocks[0][2] == "Pollination"
    assert load_source(txt, "S002").kind == "txt"
    s_dx = load_source(dx, "S003")
    assert s_dx.blocks[0][2] == "Volcanoes" and "lava" in s_dx.blocks[0][0]
    s_pdf = load_source(pdf, "S004")
    assert s_pdf.pages == 2 and s_pdf.blocks[1][1] == 2


def test_unsupported_and_missing_files(tmp_path):
    bad = tmp_path / "x.exe"
    bad.write_bytes(b"MZ")
    with pytest.raises(IngestionError):
        load_source(bad, "S1")
    with pytest.raises(IngestionError):
        load_source(tmp_path / "missing.pdf", "S1")


def test_bm25_search_keeps_source_and_page(docs):
    chunks = [c for i, p in enumerate(docs, 1) for c in chunk(load_source(p, f"S{i:03d}"))]
    idx = ChunkIndex(chunks)
    hit, _ = idx.search("how long does the moon take to orbit earth")[0]
    assert hit.source == "space.pdf" and hit.page == 2
    hit, _ = idx.search("lava magma")[0]
    assert hit.source == "geo.docx" and hit.heading == "Volcanoes"


def test_verify_claim_supported_vs_unsupported(docs):
    chunks = [c for i, p in enumerate(docs, 1) for c in chunk(load_source(p, f"S{i:03d}"))]
    idx = ChunkIndex(chunks)
    ch, ev, score = verify_claim(idx, "The Moon orbits the Earth about once every 27 days.")
    assert ch and ch.page == 2 and score >= 0.6
    # same topic, invented number -> must not be supported
    _, _, score2 = verify_claim(idx, "The Moon orbits the Earth about once every 45 days.")
    assert score2 < 0.6
    _, _, score3 = verify_claim(idx, "Penguins can fly over the Sahara desert.")
    assert score3 < 0.6


def test_numbers_supported():
    assert numbers_supported("about 27 days", "once every 27 days")
    assert not numbers_supported("about 2.7 days", "once every 27 days")
    assert numbers_supported("no numbers here", "anything")


def test_text_source_chunking():
    src = text_source("## A\n" + "word " * 300, "S9")
    ch = chunk(src, words=100, overlap=10)
    assert len(ch) >= 3 and all(c.source_id == "S9" and c.heading == "A" for c in ch)
